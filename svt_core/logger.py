import os
import sys
import time
import logging
from contextvars import ContextVar
from enum import Enum
from logging.handlers import RotatingFileHandler

from datetime import datetime
import pandas as pd
from uuid import uuid4


PASS_LEVEL = 25
logging.addLevelName(PASS_LEVEL, "PASS")

FAIL_LEVEL = 35
logging.addLevelName(FAIL_LEVEL, "FAIL")


class LogLevel(Enum):
    """Log levels mapped to Python logging levels."""

    DEBUG = (logging.DEBUG, "DEBG")
    INFO = (logging.INFO, "INFO")
    PASS = (PASS_LEVEL, "PASS")
    WARNING = (logging.WARNING, "WARN")
    FAIL = (FAIL_LEVEL, "FAIL")
    ERROR = (logging.ERROR, "ERRR")
    CRITICAL = (logging.CRITICAL, "CRIT")

    def __init__(self, level: int, display: str):
        self.level = level
        self.display = display


class SVTFormatter(logging.Formatter):
    """Custom formatter matching the original log format."""

    def format(self, record: logging.LogRecord) -> str:
        timestamp = time.strftime("%H:%M:%S")
        level_display = {
            logging.DEBUG: "DEBG",
            logging.INFO: "INFO",
            PASS_LEVEL: "PASS",
            logging.WARNING: "WARN",
            FAIL_LEVEL: "FAIL",
            logging.ERROR: "ERRR",
            logging.CRITICAL: "CRIT",
        }.get(record.levelno, "???")

        context_parts = [record.threadName or ""]
        cycle = getattr(record, "cycle", None)
        test_name = getattr(record, "test_name", None)
        if cycle is not None:
            context_parts.append(f"c={cycle}")
        if test_name is not None:
            context_parts.append(f"t={test_name}")

        prefix = f"[{timestamp}] | {level_display} | [{', '.join(context_parts)}] | "

        message = prefix + record.getMessage()

        # Append traceback if exception info is present
        if record.exc_info and not record.exc_text:
            record.exc_text = self.formatException(record.exc_info)
        if record.exc_text:
            message = message + "\n" + record.exc_text

        return message


class TitleFormatter(logging.Formatter):
    """Formatter for title messages (no timestamp/level prefix)."""

    def format(self, record: logging.LogRecord) -> str:
        return record.getMessage()


_LOG_CONTEXT: ContextVar[dict[str, object]] = ContextVar("svt_log_context", default={})


class ContextFilter(logging.Filter):
    """Injects per-thread context values into each log record."""

    def filter(self, record: logging.LogRecord) -> bool:
        context = _LOG_CONTEXT.get({})
        for key, value in context.items():
            setattr(record, key, value)
        return True


class Logger:
    """
    A logger that wraps Python's logging module.

    Example usage:
        logger = Logger.setup("device1", "stress_test")
        logger.info("Test started")
        logger.error("Something went wrong")
    """

    def __init__(
        self,
        logger: logging.Logger,
        title_logger: logging.Logger,
        log_file: str = "",
        summary_log_file: str = "",
    ):
        self._logger = logger
        self._title_logger = title_logger
        self.log_file = log_file
        self.summary_log_file = summary_log_file

    @classmethod
    def setup(
        cls,
        device_under_test: str,
        test_type: str,
        ambient: int | None = None,
        base_dir: str = "",
        extra_dirs: list[str] | None = None,
        min_level: LogLevel = LogLevel.DEBUG,
        summary_level: LogLevel = LogLevel.PASS,
        print_level: LogLevel = LogLevel.INFO,
        max_bytes: int = 200 * 1024 * 1024,
        backup_count: int = 5,
    ) -> "Logger":
        """
        Creates a Logger with a log file in a structured folder hierarchy.

        Args:
            device_under_test: Name of the device being tested
            test_type: Type of test being performed
            ambient: Ambient temperature in Celsius to include in filename
            base_dir: Primary base directory for logs
            extra_dirs: Additional directories to mirror logs into
            min_level: Minimum log level to record in the full log file
            summary_level: Minimum log level for the summary log file (default: PASS)
            print_level: Minimum log level to print to console
        Returns:
            Configured Logger instance
        """
        import inspect

        if not base_dir:
            caller_frame = inspect.stack()[1]
            caller_file = caller_frame.filename
            main_dir = os.path.dirname(os.path.abspath(caller_file))
            base_dir = os.path.join(main_dir, "logs")

        full_dir = os.path.join(base_dir, device_under_test, test_type)
        os.makedirs(full_dir, exist_ok=True)
        uuid = uuid4().hex[:8]  # Short unique identifier to avoid filename collisions
        timestamp = time.strftime("%Y%m%d-%H%M%S")
        ambient_part = f"_amb{ambient}C" if ambient is not None else ""
        log_filename = (
            f"{device_under_test}_{test_type}{ambient_part}_{timestamp}_{uuid}.log"
        )
        log_file = os.path.join(full_dir, log_filename)
        summary_log_filename = f"{device_under_test}_{test_type}{ambient_part}_{timestamp}_{uuid}_summary.log"
        summary_log_file = os.path.join(full_dir, summary_log_filename)

        # Create unique logger name to avoid conflicts
        logger_name = f"svt.{device_under_test}.{test_type}.{timestamp}.{uuid}"

        # Main logger for standard log messages
        logger = logging.getLogger(logger_name)
        logger.setLevel(min_level.level)
        logger.handlers.clear()
        logger.filters.clear()
        logger.addFilter(ContextFilter())

        formatter = SVTFormatter()

        # Console handler
        console_handler = logging.StreamHandler(sys.stdout)
        console_handler.setLevel(print_level.level)
        console_handler.setFormatter(formatter)
        logger.addHandler(console_handler)

        # File handler with rotation to avoid giant single files during long runs.
        file_handler = RotatingFileHandler(
            log_file,
            maxBytes=max_bytes,
            backupCount=backup_count,
            encoding="utf-8",
        )
        file_handler.setLevel(min_level.level)
        file_handler.setFormatter(formatter)
        logger.addHandler(file_handler)

        # Summary file handler — only records summary_level and above
        summary_file_handler = RotatingFileHandler(
            summary_log_file,
            maxBytes=max_bytes,
            backupCount=backup_count,
            encoding="utf-8",
        )
        summary_file_handler.setLevel(summary_level.level)
        summary_file_handler.setFormatter(formatter)
        logger.addHandler(summary_file_handler)

        # Mirror log to any extra directories (e.g. local project logs/)
        for extra_base in extra_dirs or []:
            try:
                extra_full_dir = os.path.join(extra_base, device_under_test, test_type)
                os.makedirs(extra_full_dir, exist_ok=True)
                extra_log_file = os.path.join(
                    extra_full_dir, os.path.basename(log_file)
                )
                extra_fh = RotatingFileHandler(
                    extra_log_file,
                    maxBytes=max_bytes,
                    backupCount=backup_count,
                    encoding="utf-8",
                )
                extra_fh.setLevel(min_level.level)
                extra_fh.setFormatter(formatter)
                logger.addHandler(extra_fh)

                extra_summary_log_file = os.path.join(
                    extra_full_dir, os.path.basename(summary_log_file)
                )
                extra_summary_fh = RotatingFileHandler(
                    extra_summary_log_file,
                    maxBytes=max_bytes,
                    backupCount=backup_count,
                    encoding="utf-8",
                )
                extra_summary_fh.setLevel(summary_level.level)
                extra_summary_fh.setFormatter(formatter)
                logger.addHandler(extra_summary_fh)
            except OSError:
                pass  # Skip extra dir if it can't be created (e.g. remote path unavailable)

        # Separate logger for titles (no formatting)
        title_logger = logging.getLogger(f"{logger_name}.title")
        title_logger.setLevel(logging.INFO)
        title_logger.handlers.clear()

        title_formatter = TitleFormatter()

        title_console = logging.StreamHandler(sys.stdout)
        title_console.setFormatter(title_formatter)
        title_logger.addHandler(title_console)

        title_file = logging.FileHandler(log_file, encoding="utf-8")
        title_file.setFormatter(title_formatter)
        title_logger.addHandler(title_file)

        summary_title_file = logging.FileHandler(summary_log_file, encoding="utf-8")
        summary_title_file.setFormatter(title_formatter)
        title_logger.addHandler(summary_title_file)

        for extra_base in extra_dirs or []:
            try:
                extra_full_dir = os.path.join(extra_base, device_under_test, test_type)
                extra_log_file = os.path.join(
                    extra_full_dir, os.path.basename(log_file)
                )
                extra_title_fh = logging.FileHandler(extra_log_file, encoding="utf-8")
                extra_title_fh.setFormatter(title_formatter)
                title_logger.addHandler(extra_title_fh)

                extra_summary_log_file = os.path.join(
                    extra_full_dir, os.path.basename(summary_log_file)
                )
                extra_summary_title_fh = logging.FileHandler(
                    extra_summary_log_file, encoding="utf-8"
                )
                extra_summary_title_fh.setFormatter(title_formatter)
                title_logger.addHandler(extra_summary_title_fh)
            except OSError:
                pass

        # Prevent propagation to root logger
        logger.propagate = False
        title_logger.propagate = False

        return cls(
            logger=logger,
            title_logger=title_logger,
            log_file=log_file,
            summary_log_file=summary_log_file,
        )

    def debug(self, msg: str) -> None:
        """Log a debug message."""
        self._logger.debug(msg)

    def info(self, msg: str) -> None:
        """Log an info message."""
        self._logger.info(msg)

    def warning(self, msg: str) -> None:
        """Log a warning message."""
        self._logger.warning(msg)

    def error(self, msg: str) -> None:
        """Log an error message."""
        self._logger.error(msg)

    def critical(self, msg: str) -> None:
        """Log a critical message."""
        self._logger.critical(msg)

    def exception(self, msg: str) -> None:
        """Log an error message with traceback information."""
        self._logger.exception(msg)

    def bind_context(self, **kwargs: object) -> None:
        """Bind per-thread contextual fields to be included in each log record."""
        current = dict(_LOG_CONTEXT.get({}))
        current.update({k: v for k, v in kwargs.items() if v is not None})
        _LOG_CONTEXT.set(current)

    def clear_context(self, *keys: str) -> None:
        """Clear context keys from the current thread context."""
        if not keys:
            _LOG_CONTEXT.set({})
            return
        current = dict(_LOG_CONTEXT.get({}))
        for key in keys:
            current.pop(key, None)
        _LOG_CONTEXT.set(current)

    def title(self, msg: str) -> None:
        """
        Log a title with decorative borders.

        Args:
            msg: Title text
        """
        title_block = f"\n======== {msg} ========\n"
        self._title_logger.info(title_block)

    def section(self, msg: str, char: str = "-", width: int = 60) -> None:
        """
        Log a section divider.

        Args:
            msg: Section header text
            char: Character to use for the divider line
            width: Total width of the divider
        """
        divider = char * width
        section_block = f"\n{divider}\n{msg}\n{divider}\n"
        self._title_logger.info(section_block)

    def passed(self, test_name: str) -> None:
        """Log a test pass result."""
        self._logger.log(PASS_LEVEL, f"[PASS] {test_name}")

    def failed(self, test_name: str, reason: str = "") -> None:
        """Log a test fail result."""
        msg = f"[FAIL] {test_name}"
        if reason:
            msg += f" -- {reason}"
        self._logger.log(FAIL_LEVEL, msg)


class ExcelLog:
    """
    Two-tab Excel logger. Main tests write to the "Main" sheet; background
    tests write to the "Background" sheet. Call set_thread_tab("Background")
    at the start of any background thread to route its data to the correct sheet.
    """

    def __init__(
        self,
        dut: str,
        test_type: str,
        ambient: int | None = None,
        base_dir: str = ".",
    ) -> None:
        import threading

        self.dut = dut
        uuid = uuid4().hex[:8]
        timestamp = time.strftime("%Y%m%d-%H%M%S")
        ambient_part = f"_amb{ambient}C" if ambient is not None else ""
        self.id = f"{dut}_{test_type}{ambient_part}_{timestamp}_{uuid}"
        full_dir = os.path.join(base_dir, dut, f"{test_type}_excel")
        os.makedirs(full_dir, exist_ok=True)
        self.filename = os.path.join(full_dir, f"{self.id}.xlsx")

        self._tabs: dict[str, dict[int, dict[str, str]]] = {
            "Main": {},
            "Background": {},
        }
        self._thread_local = threading.local()
        self._lock = threading.Lock()

    def _current_tab(self) -> str:
        return getattr(self._thread_local, "tab", "Main")

    def set_thread_tab(self, tab_name: str) -> None:
        """Route add_value / load_line calls from this thread to the given sheet."""
        self._thread_local.tab = tab_name

    def add_value(self, title: str, index: int, value: str) -> None:
        tab = self._current_tab()
        with self._lock:
            self._tabs[tab].setdefault(index, {})[title] = value

    def discard_row(self, index: int) -> None:
        """Remove a row from the current thread's tab without saving. Used to drop pre-test data."""
        tab = self._current_tab()
        with self._lock:
            self._tabs[tab].pop(index, None)

    def load_line(self, index: int) -> None:
        tab = self._current_tab()
        with self._lock:
            if not any(self._tabs.values()):
                return
            self._tabs[tab].setdefault(index, {})["Time"] = str(datetime.now())
            self._write()

    def _write(self) -> None:
        """Persist all non-empty tabs as separate sheets. Must be called with self._lock held."""
        import openpyxl

        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # remove the default empty sheet

        for sheet_name, data in self._tabs.items():
            if not data:
                continue
            ws = wb.create_sheet(title=sheet_name)
            df = pd.DataFrame.from_dict(data, orient="index").sort_index()

            cols = df.columns.tolist()
            if "Time" in cols:
                cols.insert(0, cols.pop(cols.index("Time")))
                df = df[cols]

            ws.append(df.columns.tolist())
            for row in df.itertuples(index=False, name=None):
                ws.append(["" if pd.isna(v) else v for v in row])

        if wb.sheetnames:
            wb.save(self.filename)
